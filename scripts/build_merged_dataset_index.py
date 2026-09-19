#!/usr/bin/env python3
"""Build the runtime spatial index for the new merged dataset source workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import sqlite3
import time
from pathlib import Path
from typing import Any, Iterable

try:
    import openpyxl
except ImportError as error:  # pragma: no cover
    raise SystemExit(
        "Missing openpyxl. Install it with: python3 -m pip install openpyxl"
    ) from error

from build_heat_index import (
    GLOBAL_RE,
    MEMBERSHIP_MAX_SPAN,
    NON_EARTH_RE,
    ROOT,
    bounds_intersect,
    load_region_catalog,
    point_in_region,
    resolve_regions,
)


DEFAULT_SOURCE = ROOT / "data/source/dataset_merged_20260909.xlsx"
DEFAULT_OUTPUT = ROOT / "data/generated/dataset_merged.sqlite"
DATA_SHEETS = ("海纳数据集", "OneEarth数据集")

COORDINATE_COLUMNS = (
    "中心经度_WGS84", "中心纬度_WGS84",
    "最小经度_WGS84", "最小纬度_WGS84",
    "最大经度_WGS84", "最大纬度_WGS84",
)
FALLBACK_CENTER_COLUMNS = ("center_lon", "center_lat")


def numeric_cell(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number != number:  # NaN
        return None
    return number


def parse_tag_list(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    parsed: Any = None
    if text.startswith("["):
        try:
            parsed = json.loads(text)
        except (TypeError, ValueError):
            parsed = None
    values = parsed if isinstance(parsed, list) else [
        item.strip(" [](){}'\"\t\r\n") for item in text.split(",")
    ]
    result: list[str] = []
    for item in values:
        tag = str(item or "").strip(" [](){}'\"\t\r\n")[:80]
        if tag and tag.casefold() not in {existing.casefold() for existing in result}:
            result.append(tag)
    return result


def dataset_key(source: str, raw_id: str, name: str) -> str:
    value = f"{source}|{raw_id}|{name}".casefold()
    return hashlib.sha1(value.encode("utf-8")).hexdigest()


def initialize_database(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        PRAGMA journal_mode=OFF;
        PRAGMA synchronous=OFF;
        PRAGMA temp_store=MEMORY;
        CREATE TABLE datasets (
          id INTEGER PRIMARY KEY,
          dataset_key TEXT NOT NULL UNIQUE,
          source TEXT NOT NULL,
          raw_id TEXT NOT NULL,
          name TEXT NOT NULL,
          name_en TEXT NOT NULL,
          description TEXT NOT NULL,
          description_en TEXT NOT NULL,
          area_range TEXT NOT NULL,
          source_area_range TEXT NOT NULL,
          tags_text TEXT NOT NULL,
          mapping_tags_text TEXT NOT NULL,
          dataset_classify TEXT NOT NULL,
          data_origin TEXT NOT NULL,
          data_origin_name TEXT NOT NULL,
          source_category TEXT NOT NULL,
          source_subcategory TEXT NOT NULL,
          data_format TEXT NOT NULL,
          crs TEXT NOT NULL,
          coordinate_status TEXT NOT NULL,
          coordinate_system TEXT NOT NULL,
          center_lon REAL,
          center_lat REAL,
          min_lon REAL,
          min_lat REAL,
          max_lon REAL,
          max_lat REAL,
          coordinate_extract_source TEXT NOT NULL,
          coordinate_validation TEXT NOT NULL,
          coordinate_note TEXT NOT NULL,
          data_access_mode TEXT NOT NULL,
          data_size TEXT NOT NULL,
          upload_time TEXT NOT NULL,
          real_upload_time TEXT NOT NULL,
          update_time TEXT NOT NULL,
          timeline_time TEXT NOT NULL,
          timeline_type TEXT NOT NULL,
          payload_json TEXT NOT NULL,
          start_year INTEGER,
          end_year INTEGER
        );
        CREATE TABLE geometries (
          id INTEGER PRIMARY KEY,
          dataset_id INTEGER NOT NULL,
          kind TEXT NOT NULL,
          west REAL NOT NULL,
          south REAL NOT NULL,
          east REAL NOT NULL,
          north REAL NOT NULL,
          region_code TEXT NOT NULL DEFAULT '',
          UNIQUE(dataset_id, kind, west, south, east, north, region_code)
        );
        CREATE TABLE region_catalog (
          region_code TEXT PRIMARY KEY,
          parent_code TEXT NOT NULL,
          level TEXT NOT NULL,
          name TEXT NOT NULL
        );
        CREATE TABLE region_memberships (
          region_code TEXT NOT NULL,
          dataset_id INTEGER NOT NULL,
          match_kind TEXT NOT NULL,
          PRIMARY KEY(region_code, dataset_id)
        );
        CREATE TABLE dataset_topics (
          dataset_id INTEGER NOT NULL,
          topic TEXT NOT NULL,
          PRIMARY KEY(dataset_id, topic)
        );
        CREATE TABLE dataset_spatial_profile (
          dataset_id INTEGER PRIMARY KEY,
          spatial_class TEXT NOT NULL
        );
        CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
        """
    )


def add_membership(memberships: dict[tuple[str, int], str], code: str, dataset_id: int, kind: str) -> None:
    if not code or code == "GLOBAL":
        return
    key = (code, dataset_id)
    existing = memberships.get(key)
    if existing is None or (existing != "point" and kind == "point"):
        memberships[key] = kind


def resolve_region_text(row: dict[str, Any]) -> str:
    # dataset_classify 是分类标签（如“全球数据”“城市数据”），不是空间覆盖，不参与区域匹配。
    return " | ".join(
        text for text in (
            str(row.get("source_area_range") or "").strip(),
            str(row.get("area_range") or "").strip(),
        )
        if text
    )


def build_index(source: Path, output: Path) -> dict[str, int | float]:
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        output.unlink()
    region_catalog, aliases = load_region_catalog()
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    connection = sqlite3.connect(output)
    initialize_database(connection)

    stats: dict[str, int | float] = {
        "rows": 0, "datasets": 0, "coordinate_rows": 0, "coordinate_geometries": 0,
        "region_rows": 0, "region_geometries": 0, "global_rows": 0, "unresolved_rows": 0,
        "non_earth_rows": 0, "topic_links": 0, "region_memberships": 0, "regions": 0,
    }
    memberships: dict[tuple[str, int], str] = {}
    started = time.perf_counter()

    try:
        with connection:
            for sheet_name in DATA_SHEETS:
                if sheet_name not in workbook.sheetnames:
                    continue
                worksheet = workbook[sheet_name]
                iterator = worksheet.iter_rows(values_only=True)
                try:
                    headers = next(iterator)
                except StopIteration:
                    continue
                header_index = {str(header): index for index, header in enumerate(headers)}

                for values in iterator:
                    stats["rows"] += 1
                    row: dict[str, Any] = {
                        str(header): values[header_index[str(header)]]
                        if header_index[str(header)] < len(values) else None
                        for header in headers
                    }
                    name = str(row.get("datanet_name") or "").strip()
                    raw_id = str(row.get("id") if row.get("id") not in (None, "") else row.get("datanet_id") or "")
                    name_en = str(row.get("datanet_name_en") or "").strip()
                    description = str(row.get("datanet_description") or "").strip()
                    description_en = str(row.get("datanet_description_en") or "").strip()
                    area_range = str(row.get("area_range") or "").strip()
                    source_area_range = str(row.get("source_area_range") or "").strip()
                    dataset_classify = str(row.get("dataset_classify") or "").strip()
                    tags = parse_tag_list(row.get("tags"))
                    mapping_tags = parse_tag_list(row.get("mapping_tags"))
                    display_tags = tags if tags else mapping_tags
                    tags_text = json.dumps(display_tags, ensure_ascii=False)
                    mapping_tags_text = json.dumps(mapping_tags, ensure_ascii=False)
                    coordinate_status = str(row.get("坐标状态") or "").strip()
                    coordinate_system = str(row.get("输出坐标系") or "").strip()
                    coordinate_extract_source = str(row.get("坐标提取来源") or "").strip()
                    coordinate_validation = str(row.get("坐标校验") or "").strip()
                    coordinate_note = str(row.get("坐标说明") or "").strip()

                    center_lon = numeric_cell(row.get("中心经度_WGS84"))
                    center_lat = numeric_cell(row.get("中心纬度_WGS84"))
                    if center_lon is None or center_lat is None:
                        center_lon = numeric_cell(row.get("center_lon"))
                        center_lat = numeric_cell(row.get("center_lat"))
                    min_lon = numeric_cell(row.get("最小经度_WGS84"))
                    min_lat = numeric_cell(row.get("最小纬度_WGS84"))
                    max_lon = numeric_cell(row.get("最大经度_WGS84"))
                    max_lat = numeric_cell(row.get("最大纬度_WGS84"))

                    payload = {
                        str(header): (values[header_index[str(header)]]
                                      if header_index[str(header)] < len(values) else None)
                        for header in headers
                    }
                    payload_json = json.dumps(payload, ensure_ascii=False, default=str)

                    key = dataset_key(sheet_name, raw_id, name)
                    cursor = connection.execute(
                        "INSERT INTO datasets("
                        "dataset_key,source,raw_id,name,name_en,description,description_en,"
                        "area_range,source_area_range,tags_text,mapping_tags_text,dataset_classify,"
                        "data_origin,data_origin_name,source_category,source_subcategory,data_format,crs,coordinate_status,coordinate_system,"
                        "center_lon,center_lat,min_lon,min_lat,max_lon,max_lat,"
                        "coordinate_extract_source,coordinate_validation,coordinate_note,"
                        "data_access_mode,data_size,upload_time,real_upload_time,update_time,"
                        "timeline_time,timeline_type,payload_json,start_year,end_year"
                        ") VALUES(" + ",".join("?" for _ in range(39)) + ")",
                        (
                            key, sheet_name, raw_id, name, name_en, description, description_en,
                            area_range, source_area_range, tags_text, mapping_tags_text, dataset_classify,
                            str(row.get("data_origin") or "").strip(),
                            str(row.get("data_origin_name") or "").strip(),
                            str(row.get("来源大类") or "").strip(),
                            str(row.get("来源子类") or "").strip(),
                            str(row.get("oridata_format") or "").strip(),
                            str(row.get("oridata_crs") or "").strip(),
                            coordinate_status, coordinate_system,
                            center_lon, center_lat, min_lon, min_lat, max_lon, max_lat,
                            coordinate_extract_source, coordinate_validation, coordinate_note,
                            str(row.get("data_access_mode") or "").strip(),
                            str(row.get("data_size") or "").strip(),
                            str(row.get("oridata_upload_time") or "").strip(),
                            str(row.get("real_upload_time") or "").strip(),
                            str(row.get("datanet_update_time") or "").strip(),
                            str(row.get("timeline_time") or "").strip(),
                            str(row.get("timeline_type") or "").strip(),
                            payload_json, None, None,
                        ),
                    )
                    dataset_id = int(cursor.lastrowid)
                    stats["datasets"] += 1

                    for topic in display_tags:
                        connection.execute(
                            "INSERT OR IGNORE INTO dataset_topics(dataset_id,topic) VALUES(?,?)",
                            (dataset_id, topic),
                        )

                    region_text = resolve_region_text(row)
                    if NON_EARTH_RE.search(coordinate_system + " " + region_text + " " + description):
                        stats["non_earth_rows"] += 1
                        continue

                    if GLOBAL_RE.search(region_text):
                        stats["region_rows"] += 1
                        stats["global_rows"] += 1
                        before = connection.total_changes
                        connection.execute(
                            "INSERT OR IGNORE INTO geometries(dataset_id,kind,west,south,east,north,region_code) "
                            "VALUES(?,'global',-180,-90,180,90,'GLOBAL')",
                            (dataset_id,),
                        )
                        stats["region_geometries"] += connection.total_changes - before
                        continue

                    has_bbox = (
                        min_lon is not None and min_lat is not None
                        and max_lon is not None and max_lat is not None
                        and -180 <= min_lon <= max_lon <= 180
                        and -90 <= min_lat <= max_lat <= 90
                    )
                    has_point = (
                        center_lon is not None and center_lat is not None
                        and -180 <= center_lon <= 180 and -90 <= center_lat <= 90
                    )
                    if has_bbox:
                        stats["coordinate_rows"] += 1
                        before = connection.total_changes
                        connection.execute(
                            "INSERT OR IGNORE INTO geometries(dataset_id,kind,west,south,east,north,region_code) "
                            "VALUES(?,?,?,?,?,?, '')",
                            (dataset_id, "bbox", min_lon, min_lat, max_lon, max_lat),
                        )
                        stats["coordinate_geometries"] += connection.total_changes - before
                        continue
                    if has_point:
                        stats["coordinate_rows"] += 1
                        before = connection.total_changes
                        connection.execute(
                            "INSERT OR IGNORE INTO geometries(dataset_id,kind,west,south,east,north,region_code) "
                            "VALUES(?,?,?,?,?,?, '')",
                            (dataset_id, "point", center_lon, center_lat, center_lon, center_lat),
                        )
                        stats["coordinate_geometries"] += connection.total_changes - before
                        continue

                    region_codes = resolve_regions(region_text, aliases)
                    if not region_codes:
                        stats["unresolved_rows"] += 1
                        continue
                    stats["region_rows"] += 1
                    for code in region_codes:
                        geometry = ("region", *region_catalog[code]["bounds"], code)
                        before = connection.total_changes
                        connection.execute(
                            "INSERT OR IGNORE INTO geometries(dataset_id,kind,west,south,east,north,region_code) "
                            "VALUES(?,?,?,?,?,?,?)",
                            (dataset_id, *geometry),
                        )
                        stats["region_geometries"] += connection.total_changes - before

            connection.executemany(
                "INSERT INTO region_catalog(region_code,parent_code,level,name) VALUES(?,?,?,?)",
                [
                    (code, str(region["parent_code"]), str(region["level"]), str(region["name"]))
                    for code, region in region_catalog.items()
                ],
            )
            coordinate_rows = connection.execute(
                "SELECT dataset_id,kind,west,south,east,north,region_code FROM geometries WHERE kind!='global'"
            ).fetchall()
            global_dataset_ids = {
                int(row[0]) for row in connection.execute(
                    "SELECT dataset_id FROM geometries WHERE kind='global'"
                ).fetchall()
            }
            regions_by_level = {
                level: [region for region in region_catalog.values() if region["level"] == level]
                for level in MEMBERSHIP_MAX_SPAN
            }
            for dataset_id, kind, west, south, east, north, region_code in coordinate_rows:
                dataset_id = int(dataset_id)
                if dataset_id in global_dataset_ids:
                    continue
                if kind == "region":
                    add_membership(memberships, str(region_code), dataset_id, "name")
                    continue
                if kind == "point":
                    longitude = (float(west) + float(east)) * 0.5
                    latitude = (float(south) + float(north)) * 0.5
                    for regions in regions_by_level.values():
                        for region in regions:
                            if point_in_region(longitude, latitude, region):
                                add_membership(memberships, region["code"], dataset_id, "point")
                    continue
                geometry_extent = (float(west), float(south), float(east), float(north))
                maximum_span = max(float(east) - float(west), float(north) - float(south))
                for level, maximum_allowed in MEMBERSHIP_MAX_SPAN.items():
                    if maximum_span > maximum_allowed:
                        continue
                    for region in regions_by_level[level]:
                        if bounds_intersect(geometry_extent, region["bounds"]):
                            add_membership(memberships, region["code"], dataset_id, "bbox")

            connection.executemany(
                "INSERT INTO region_memberships(region_code,dataset_id,match_kind) VALUES(?,?,?)",
                [(code, dataset_id, kind) for (code, dataset_id), kind in memberships.items()],
            )
            stats["region_memberships"] = len(memberships)
            stats["regions"] = len(region_catalog)
            connection.execute(
                "INSERT INTO dataset_spatial_profile(dataset_id,spatial_class) "
                "SELECT d.id,CASE "
                "WHEN MAX(CASE WHEN g.kind='global' THEN 1 ELSE 0 END)=1 THEN 'global' "
                "WHEN MAX(CASE WHEN g.kind='bbox' THEN 1 ELSE 0 END)=1 THEN 'bbox' "
                "WHEN MAX(CASE WHEN g.kind='point' THEN 1 ELSE 0 END)=1 THEN 'point' "
                "WHEN MAX(CASE WHEN g.kind='region' THEN 1 ELSE 0 END)=1 THEN 'region' "
                "ELSE 'unlocated' END "
                "FROM datasets d LEFT JOIN geometries g ON g.dataset_id=d.id GROUP BY d.id"
            )
            stats["topic_links"] = int(connection.execute(
                "SELECT COUNT(*) FROM dataset_topics"
            ).fetchone()[0])

            connection.executemany(
                "INSERT INTO metadata(key,value) VALUES(?,?)",
                [(key, str(value)) for key, value in stats.items()],
            )
            connection.execute("CREATE INDEX geometries_bounds_idx ON geometries(west,east,south,north)")
            connection.execute("CREATE INDEX geometries_east_bounds_idx ON geometries(east,west,south,north)")
            connection.execute("CREATE INDEX geometries_dataset_idx ON geometries(dataset_id)")
            connection.execute("CREATE INDEX region_memberships_region_idx ON region_memberships(region_code)")
            connection.execute("CREATE INDEX region_memberships_dataset_idx ON region_memberships(dataset_id)")
            connection.execute("CREATE INDEX dataset_topics_topic_idx ON dataset_topics(topic)")
            connection.execute("CREATE INDEX dataset_spatial_class_idx ON dataset_spatial_profile(spatial_class)")
            connection.execute("ANALYZE")
    finally:
        workbook.close()
    connection.execute("VACUUM")
    connection.close()
    stats["seconds"] = round(time.perf_counter() - started, 3)
    return stats


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    stats = build_index(args.source.resolve(), args.output.resolve())
    print(json.dumps(stats, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
